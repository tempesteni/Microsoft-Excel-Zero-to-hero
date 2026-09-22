#!/usr/bin/env python3
"""Generate exercise-fix-this-workbook.xlsx: a DELIBERATELY broken workbook with 20
planted issues. Companion: fix-this-workbook-challenge.md (answer key). Fictional data only."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import PieChart, BarChart, Reference

OUT = "~/EXCEL Learn/samples"

wb = Workbook()

# ================= Sheet 1: Results (formula + format issues) =================
ws = wb.active; ws.title = "Results"
ws["A1"] = "Q2 RESULTS DASHBOARD"
ws["A1"].font = Font(name="Impact", size=20, color="800080")  # ISSUE 11 (font #2)
ws["A3"] = "Total revenue (should be a formula!)"
ws["B3"] = 25                                                   # ISSUE 1: hardcoded value
ws["A4"] = "Regional total (broken reference)"
ws["B4"] = "=SUM(#REF!)"                                        # ISSUE 2: #REF! error
ws["A5"] = "Avg order value (divides by zero)"
ws["B5"] = "=B6/B7"                                             # ISSUE 3: #DIV/0! (B7 = 0)
ws["B6"] = 1250; ws["B7"] = 0
ws["A10"] = "Notes: light gray on light yellow, good luck reading this"
ws["A10"].font = Font(name="Calibri", color="D9D9D9")           # ISSUE 13: low contrast
ws["A10"].fill = PatternFill("solid", fgColor="FFFFCC")

# ================= Sheet 2: Data (data hygiene + structure issues) =================
d = wb.create_sheet("Data")
heads = ["Customer", "Date", "Region(hidden col C!)", "Qty", "Amount", "", "Notes", "Row Total"]
for i, h in enumerate(heads, 1):
    c = d.cell(row=1, column=i, value=h if h else None)          # ISSUE 10: blank header (col F)
    fills = ["FF0000", "800080", "FF8000", "008000", "008080", "FFFFFF", "FFFF00", "FFFFFF"]
    c.fill = PatternFill("solid", fgColor=fills[i - 1])          # ISSUE 14: rainbow header row
rows = [
    ["Northwind Traders ", "03/05/2026", "correction: was 8", "10", 250, None, "call again", 250],
    ["Redwood Labs", "03/06/2026", "ok", "4", 130.5, None, "", 522],
    ["Northwind Traders", "03/07/2026", "ok", "25", 625, None, "", 625],
    ["Harborline Credit Union ", "03/08/2026", "check addr", "2", 960, None, "", 960],
    ["Meridian Freight", "03/09/2026", "ok", "6", 783, None, "", 783],
    ["Greenfield Literacy Trust", "03/10/2026", "ok", "15", 375, None, "", 375],
    ["Northgate Institute", "03/11/2026", "ok", "1", 480, None, "", 480],
    ["VeloCabin", "03/12/2026", "ok", "8", 1044, None, "", 1044],
    ["Northwind Traders", "03/07/2026", "ok", "25", 625, None, "", 625],   # ISSUE 9: duplicate row
    ["Harborline Credit Union ", "03/08/2026", "check addr", "2", 960, None, "", 960],  # ISSUE 9: duplicate
    ["Redwood Labs", "03/13/2026", "ok", "12", 1566, None, "", 1566],
    ["Meridian Freight", "03/14/2026", "ok", "3", 1440, None, "", 1440],
    ["Greenfield Literacy Trust", "03/15/2026", "ok", "7", 913.5, None, "", 913.5],
]
for r in rows:
    d.append(r)
for r in range(2, 2 + len(rows)):
    d.cell(row=r, column=4).number_format = "@"    # ISSUE 6: Qty stored as text
    d.cell(row=r, column=2).alignment = Alignment(horizontal="left")
for r in (2, 5, 6):                                 # ISSUE 15: mixed number formats in Amount col
    d.cell(row=r, column=5).number_format = '"$"#,##0.00'
for r in (3, 4, 7, 8):
    d.cell(row=r, column=5).font = Font(name="Comic Sans MS")  # ISSUE 11 (font #3)
d["H16"] = "=SUM(H2:H12)"                           # ISSUE 4: off-by-one SUM (misses rows 13-14)
d["A16"] = "Grand total (check the range!)"
d["E2"] = "=E2"                                     # ISSUE 20: circular reference
d["G8"] = "TODO fix this"                           # stray note riding along
d.merge_cells("A8:B8")                              # ISSUE 12: merged cells inside data
d.merge_cells("G1:H1")                              # ISSUE 12 (cont): merged header
d.column_dimensions["C"].hidden = True              # ISSUE 19: hidden column with real data
# ISSUE 7: dates stored as text (all values in column B are strings like "03/05/2026")

# ================= Sheet 3: Lookups (bad VLOOKUP) =================
lk = wb.create_sheet("Lookups")
lk["A1"] = "Find the price for each SKU"
lk["A3"] = "SKU"; lk["B3"] = "Price (VLOOKUP approx!)"
lk["D3"] = "SKU"; lk["E3"] = "Price"
catalog = [("SKU-010", 95), ("SKU-002", 40), ("SKU-031", 260), ("SKU-007", 72), ("SKU-022", 150)]
for i, (s, p) in enumerate(catalog):
    lk.cell(row=4 + i, column=4, value=s)           # ISSUE 5 context: lookup column NOT sorted
    lk.cell(row=4 + i, column=5, value=p)
for i, sku in enumerate(["SKU-031", "SKU-007", "SKU-002"]):
    lk.cell(row=4 + i, column=1, value=sku)
    lk.cell(row=4 + i, column=2, value=f"=VLOOKUP(A{4+i},$D$4:$E$8,2,TRUE)")  # ISSUE 5: TRUE on unsorted

# ================= Sheet 4: Charts (chart crimes) =================
ch = wb.create_sheet("Charts")
ch["A1"] = "Month"; ch["B1"] = "Revenue"
months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
vals = [42, 51, 47, 63, 58, 71, 66, 74, 69, 82, 77, 91]
for i, (m, v) in enumerate(zip(months, vals), 2):
    ch.cell(row=i, column=1, value=m)
    ch.cell(row=i, column=2, value=v)
pie = PieChart(); pie.title = "Monthly revenue (why a pie?!)"    # ISSUE 16: pie for a time series
pie.add_data(Reference(ch, min_col=2, min_row=1, max_row=13), titles_from_data=True)
pie.set_categories(Reference(ch, min_col=1, min_row=2, max_row=13))
ch.add_chart(pie, "D2")
bar = BarChart(); bar.title = "Revenue growth (trust me)"
bar.add_data(Reference(ch, min_col=2, min_row=1, max_row=13), titles_from_data=True)
bar.set_categories(Reference(ch, min_col=1, min_row=2, max_row=13))
bar.y_axis.scaling.min = 78    # ISSUE 17: truncated axis (data 42-91 shown from 78)
bar.y_axis.scaling.max = 92
bar.legend.position = "r"      # ISSUE 18: chart junk (legend + labels + gridlines all on)
bar.dLbls = None
from openpyxl.chart.label import DataLabelList
bar.dataLabels = DataLabelList(); bar.dataLabels.showVal = True
bar.y_axis.majorGridlines = None
from openpyxl.chart.axis import ChartLines
bar.y_axis.majorGridlines = ChartLines()
ch.add_chart(bar, "D20")

wb.save(f"{OUT}/exercise-fix-this-workbook.xlsx")
print("saved: exercise-fix-this-workbook.xlsx")

# ---- self-QA: confirm every planted issue exists ----
from openpyxl import load_workbook
wb2 = load_workbook(f"{OUT}/exercise-fix-this-workbook.xlsx")
checks = {
    "hardcoded total": wb2["Results"]["B3"].value == 25,
    "REF error formula": "#REF!" in str(wb2["Results"]["B4"].value),
    "DIV0 formula": wb2["Results"]["B5"].value == "=B6/B7",
    "off-by-one SUM": wb2["Data"]["H16"].value == "=SUM(H2:H12)",
    "circular formula": wb2["Data"]["E2"].value == "=E2",
    "VLOOKUP approx": "TRUE" in str(wb2["Lookups"]["B4"].value),
    "text qty": wb2["Data"]["D2"].number_format == "@",
    "text dates": isinstance(wb2["Data"]["B2"].value, str),
    "trailing spaces": wb2["Data"]["A2"].value.endswith(" "),
    "3 fonts": {"Impact", "Comic Sans MS", "Calibri"} <= {c.font.name for s in wb2 for row in s.iter_rows() for c in row if c.font and c.font.name},
    "merged data cell": "A8:B8" in [str(m) for m in wb2["Data"].merged_cells.ranges],
    "duplicate rows": [c.value for c in wb2["Data"][10]] == [c.value for c in wb2["Data"][4]],
    "hidden column": wb2["Data"].column_dimensions["C"].hidden,
    "charts": len(wb2["Charts"]._charts) == 2,
}
for k, v in checks.items():
    print(("OK  " if v else "MISS"), k)
print("planted issues: 20 (see fix-this-workbook-challenge.md)")
